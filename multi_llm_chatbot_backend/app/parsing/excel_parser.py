"""Excel (.xlsx) parser — every sheet as pipe-delimited rows."""

from __future__ import annotations

from io import BytesIO

from app.parsing.base import BaseParser


class ExcelParser(BaseParser):
    """Extract each worksheet's rows from an .xlsx workbook."""

    def parse(self, file_bytes: bytes) -> str:
        from openpyxl import load_workbook  # type: ignore

        workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                if any(c.strip() for c in cells):
                    parts.append(" | ".join(cells))
            parts.append("")
        return "\n".join(parts).strip()
