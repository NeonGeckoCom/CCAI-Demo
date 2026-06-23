"""Word document parser.

Uses ``python-docx`` so tables are captured (TrustRAG's ``docx_parser.py``
does the same). Falls back to ``docx2txt`` if python-docx fails.
"""

from __future__ import annotations

import logging
import os
import tempfile
from io import BytesIO

import docx2txt

from app.parsing.base import BaseParser

logger = logging.getLogger(__name__)


class DocxParser(BaseParser):
    """Extract paragraphs and tables from a .docx file."""

    def parse(self, file_bytes: bytes) -> str:
        try:
            from docx import Document  # python-docx

            doc = Document(BytesIO(file_bytes))
            parts = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    parts.append(text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as e:
            logger.warning(f"python-docx failed ({e}); falling back to docx2txt")
            with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name
            try:
                return docx2txt.process(tmp_path)
            finally:
                os.unlink(tmp_path)
