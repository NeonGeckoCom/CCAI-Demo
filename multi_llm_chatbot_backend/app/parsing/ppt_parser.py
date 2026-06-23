"""PowerPoint (.pptx) parser — slide text, tables, and speaker notes."""

from __future__ import annotations

from io import BytesIO

from app.parsing.base import BaseParser


class PptParser(BaseParser):
    """Extract slide text, table cells, and speaker notes from a .pptx deck."""

    def parse(self, file_bytes: bytes) -> str:
        from pptx import Presentation  # type: ignore

        prs = Presentation(BytesIO(file_bytes))
        parts = []
        for idx, slide in enumerate(prs.slides, start=1):
            parts.append(f"# Slide {idx}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs).strip()
                        if text:
                            parts.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"[Speaker notes] {notes}")
            parts.append("")
        return "\n".join(parts).strip()
